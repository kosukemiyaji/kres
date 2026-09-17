"""
data/import_addresses.py
-------------------------
総務省が公開している「全国地方公共団体コード」のExcel
（シート1: R6.1.1現在の団体 / シート2: R6.1.1政令指定都市）
を読み込み、db/schema.sql の addresses テーブルへ投入する。

シート1: 全都道府県・市区町村（都道府県のみの行は市区町村名が空）
シート2: 政令指定都市とその区の一覧（区かどうかの判定に使う）

実行方法:
    python -m data.import_addresses /path/to/000925835.xlsx
"""

import sys
import sqlite3
import os
import openpyxl
from pathlib import Path

# bot.py と同じ DATABASE_PATH を参照するため、Render の永続ディスクへ
# 初期データを投入する場合にも同じ DB を対象にできる。
DB_PATH = Path(os.environ.get("DATABASE_PATH", Path(__file__).parent.parent / "db" / "econobot.db"))
SCHEMA_PATH = Path(__file__).parent.parent / "db" / "schema.sql"


def load_workbook_rows(xlsx_path: str):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    main_sheet = wb["R6.1.1現在の団体"]
    designated_sheet = wb["R6.1.1政令指定都市"]

    main_rows = list(main_sheet.iter_rows(min_row=2, values_only=True))
    designated_rows = list(designated_sheet.iter_rows(min_row=2, values_only=True))

    return main_rows, designated_rows


def _upsert(conn, code, pref_name, city_name, pref_kana, city_kana, is_ward):
    code = str(code).zfill(6)
    pref_code = code[:2]
    conn.execute(
        "INSERT INTO addresses (code, pref_code, pref_name, pref_kana, city_name, city_kana, "
        "is_designated_city_ward) VALUES (?, ?, ?, ?, ?, ?, ?) "
        "ON CONFLICT(code) DO UPDATE SET "
        "pref_name=excluded.pref_name, city_name=excluded.city_name, "
        "pref_kana=excluded.pref_kana, city_kana=excluded.city_kana, "
        "is_designated_city_ward=excluded.is_designated_city_ward",
        (code, pref_code, pref_name, pref_kana, city_name, city_kana, is_ward),
    )


def import_addresses(xlsx_path: str):
    main_rows, designated_rows = load_workbook_rows(xlsx_path)

    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    with open(SCHEMA_PATH, encoding="utf-8") as f:
        conn.executescript(f.read())

    inserted = 0

    # 1) メインシート：都道府県 + 通常の市区町村 + 東京都特別区
    for row in main_rows:
        code, pref_name, city_name, pref_kana, city_kana = row[0], row[1], row[2], row[3], row[4]
        if not code or not pref_name:
            continue
        _upsert(conn, code, pref_name, city_name, pref_kana, city_kana, is_ward=0)
        inserted += 1

    # 2) 政令指定都市シート：「市」の行はメインシートと同一コードで重複するため無視し、
    #    「区」の行（市区町村名に「区」を含む）だけを追加登録する。
    #    政令指定都市の区はメインシートに単独では載っていない。
    ward_inserted = 0
    for row in designated_rows:
        code, pref_name, city_name, pref_kana, city_kana = row[0], row[1], row[2], row[3], row[4]
        if not code or not city_name or "区" not in str(city_name):
            continue
        _upsert(conn, code, pref_name, city_name, pref_kana, city_kana, is_ward=1)
        ward_inserted += 1

    conn.commit()
    conn.close()
    print(f"取り込み完了: 都道府県・市区町村 {inserted} 件 + 政令指定都市の区 {ward_inserted} 件 -> {DB_PATH}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("使い方: python -m data.import_addresses <xlsxファイルパス>")
        sys.exit(1)
    import_addresses(sys.argv[1])
